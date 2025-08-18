from __future__ import annotations

from .base import BaseParser
from ...domain.models import ParsedDocument, RawDocument


class DocxParser(BaseParser):
    """DOCX -> text via python-docx (if available).

    - Concatenates paragraphs with newlines.
    - Skips non-printable control characters.
    """

    def parse(self, raw: RawDocument) -> ParsedDocument:
        text = ""
        try:
            from docx import Document  # type: ignore

            doc = Document(str(raw.path))
            parts = [p.text for p in doc.paragraphs if p.text is not None]
            text = "\n".join(parts)
        except Exception:
            text = ""
        # Remove non-printable controls except common whitespace
        def _clean(s: str) -> str:
            return "".join(ch for ch in s if (ch.isprintable() or ch in "\n\r\t"))

        cleaned = _clean(text)
        metadata = {"ext": "docx", "bytes": raw.size}
        return ParsedDocument(text=cleaned, source=raw, metadata=metadata)

